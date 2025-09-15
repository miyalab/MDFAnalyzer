import os
import asammdf
import shutil
import inspect
import threading
import collections
import openpyxl
import openpyxl.cell.cell
import time
import traceback
from glob import glob

from typing import List, Dict

class MDFAnalyzer:
    def __init__(self, xlsx_path: str):
        self.__xlsx: openpyxl.Workbook = None
        self.__xlsx_path: str = ""
        self.__analysis_data: openpyxl.cell.cell = None
        self.__mdf_dir_data: openpyxl.cell.cell.Cell = None
        self.__calc_rate_data: openpyxl.cell.cell.Cell = None
        self.__function_info_data: openpyxl.cell.cell.Cell = None
        self.__labels_data: openpyxl.cell.cell.Cell = None
        self.__loadAnalysisData(xlsx_path)

    def save(self):
        self.__xlsx.save(self.__xlsx_path)

    def updateMDFList(self):
        self.__clearAnalysisData()
        
        row: int = 3
        for path in sorted(path.replace("\\", "/") for path in glob(f"{self.__mdf_dir_data.value}*.mf4")):
            self.__analysis_data(row, 1).value = path.split("/")[-1]
            row += 1

        col : int = 2
        self.__analysis_data(1,1).value = "FileName"
        for func in self.__function_info_data.value.split(","):
            items = func.split(":")
            self.__analysis_data(1, col).value = items[0]
            self.__analysis_data(2, col).value = items[1]
            col += 1

    def updateLabelList(self, functions: List):
        labels = []
        
        for func in functions:
            labels += self.__getLabelListInFunc(inspect.getsource(func))
        self.__labels_data.value = ",".join(list(collections.Counter(labels).keys()))

    def setConfig(self, mdf_dir: str = None, calc_rate: float = None, functions: Dict = None):
        analysis_update = False

        ### MDFディレクトリ
        if mdf_dir is not None:
            if mdf_dir[-1] != "/": mdf_dir += "/"
            if self.__mdf_dir_data.value != mdf_dir:
                analysis_update = True
                self.__mdf_dir_data.value = mdf_dir
        
        ### 計算関数の変更
        if functions is not None:
            info = ",".join([f"{key}:{value.__name__}" for key, value in functions.items()])
            if self.__function_info_data.value != info:
                analysis_update = True
                self.__function_info_data.value = info
                self.updateLabelList(functions.values())

        ### 計算レート変更
        if calc_rate is not None:
            if self.__calc_rate_data.value != calc_rate:
                analysis_update = True
                self.__calc_rate_data.value = calc_rate

        if analysis_update: self.updateMDFList()

    def calculate(self, re_calc: bool = False):
        ### 全部再計算
        if re_calc: self.updateMDFList()

        ### MDFダウンロード先フォルダ
        os.makedirs("temp/", exist_ok=True)

        download_list: List[str] = []
        downloader: threading.Thread = None
        threads: List[threading.Thread] = []
        row: int = 3
        while self.__analysis_data(row,1).value is not None:
            col: int = 2
            functions: Dict[int, str] = {}
            while self.__analysis_data(1, col).value is not None:
                ### 数値以外であれば計算対象
                if type(self.__analysis_data(row, col).value) != float and type(self.__analysis_data(row, col).value) != int:
                    functions[col] = self.__analysis_data(2, col).value
                col += 1

            ### 計算対象があれば，別スレッドで計算して格納まで実施
            if len(functions.keys()) > 0:
                thread = threading.Thread(
                    target=self.__analysisMDF,
                    args=[self.__analysis_data(row, 1).value, row, functions],
                    name=self.__analysis_data(row, 1).value
                )
                threads.append(thread)
                download_list.append(self.__analysis_data(row, 1).value)
            row += 1
        
        if len(download_list) > 0:
            downloader = threading.Thread(target=self.__downloadMDF, args=[download_list])
            downloader.start()
        for thread in threads: 
            thread.start()
            thread.join()
        if downloader is not None: downloader.join()

    def __downloadMDF(self, mdf_names: List[str]):
        ### MDFをコピー
        for mdf_name in mdf_names:
            shutil.copy(self.__mdf_dir_data.value + mdf_name, f"temp/{mdf_name}")

    def __analysisMDF(self, mdf_name: str, row: int, functions: Dict[int, str]):
        ### ダウンロード完了待ち
        while not os.path.exists(f"temp/{mdf_name}"): time.sleep(1)
        
        try:
            ### MDF読み込み
            mdf = asammdf.MDF(f"temp/{mdf_name}")
            dataframe = mdf.to_dataframe(self.__labels_data.value.split(","), raster=self.__calc_rate_data.value)

            ### データフレームを渡す
            for col, func in functions.items():
                self.__callFunction(func, mdf_name, dataframe, self.__analysis_data(row, col))
        except Exception as e:
            pass
        
        ### コピーファイルを削除（失敗時は最大5回まで再試行する）
        for i in range(5):
            try: 
                os.remove(f"temp/{mdf_name}")
                break
            except PermissionError as e: time.sleep(1)

    def __callFunction(self, func: str, mdf_name: str, dataframe, result: openpyxl.cell.cell.Cell):
        result.value = eval(f"{func}(mdf_name, dataframe)")

    def __getLabelListInFunc(self, code: str) -> List[str]:
        ret: List[str] = []
        index1 = 0
        index2 = 0
        index1 = code.find("dataframe[", index2)
        while index1 >= 0:
            index2 = code.find("]", index1+10)
            ret.append(code[index1+11:index2-1])
            index1 = code.find("dataframe[\"", index2)
        return ret

    def __loadAnalysisData(self, xlsx_path: str):
        xlsx: openpyxl.Workbook = None
        if os.path.exists(xlsx_path):
            xlsx = openpyxl.load_workbook(xlsx_path)
        else:
            xlsx = openpyxl.Workbook()
            xlsx.create_sheet("setting")
            xlsx.create_sheet("analysis")
            del xlsx["Sheet"]

            xlsx["setting"].cell(1,1).value = "MDF Directory"
            xlsx["setting"].cell(2,1).value = "Calculate Rate"
            xlsx["setting"].cell(3,1).value = "Calculate Function"
            xlsx["setting"].cell(4,1).value = "Calculate Label"

        self.__xlsx = xlsx
        self.__xlsx_path = xlsx_path
        self.__mdf_dir_data = xlsx["setting"].cell(1,2)
        self.__calc_rate_data = xlsx["setting"].cell(2,2)
        self.__function_info_data = xlsx["setting"].cell(3,2)
        self.__labels_data = xlsx["setting"].cell(4,2)
        self.__analysis_data = xlsx["analysis"].cell

    def __clearAnalysisData(self):
        del self.__xlsx["analysis"]
        self.__analysis_data = self.__xlsx.create_sheet("analysis").cell
    
def calc1(filename, dataframe) -> float:
    print(f"{threading.current_thread().name}: call calc1")
    dataframe["LABEL1"]
    return 1.0

def calc2(filename, dataframe) -> float:
    print(f"{threading.current_thread().name}: call calc2")
    dataframe["LABEL2"]
    dataframe["LABEL3"]
    dataframe["LABEL4"]
    dataframe["LABEL5"]
    dataframe["LABEL6"]
    return 2.0

if __name__ == "__main__":
    app = MDFAnalyzer("analysis.xlsx")
    app.setConfig("C:/Users/xxxxx/Downloads/mdfdata", 0.1, {"計算1": calc1, "計算2": calc2})
    app.calculate()
    app.save()
